"""File S3, Range Maps of Species in Climatic Gaps: every species that is threatened (CR, EN or VU) or has > 90 % of its
global range in Peru (at sea, in the 200 nautical miles), and has > 50 % of its range in the domain in climatic gaps (at sea,
in at least one layer); rule in especies.seleccion_mapas.

  Fig. S3.1 amphibians, S3.2 mammals, S3.3 reptiles, S3.4 plants, S3.5 birds (16 per page); S3.6 marine fishes, mammals and
  reptiles (sea surface and seafloor, 8 per page); S3.7 seabirds (sea surface only, 16 per page)

Writes figuras/especies/ (pages), figuras/especies/sig/ (class rasters and range shapefiles) and
figuras/especies/File_S3_Species_range_maps.docx. All of it derives from IUCN and BirdLife ranges, which cannot be
redistributed, so it stays outside git.
Run: python mapas_especies.py [png] [sig] [docx]   (the three steps without arguments)
"""
import re
import shutil
import sys

import pyogrio
import especies as E
from figuras_comun import *  # noqa: F401,F403  (shared constants, canvases and map helpers)

DST = E.ESPECIES
DOCX = DST / 'File_S3_Species_range_maps.docx'
T, M = E.cargar()
Q, QM = E.seleccion_mapas(T, M)
QM_PECES, QM_AVES = QM[QM.grupo != 'BIRDS'], QM[QM.grupo == 'BIRDS']
GRUPOS = E.TIERRA_G
MAR_FIG, AVES_FIG = f"Fig{E.ATLAS_FIG['MARINEFISH']}_marine_species", f"Fig{E.ATLAS_FIG['SEABIRDS']}_seabirds"
pasos = sys.argv[1:] or ['png', 'sig', 'docx']
assert len(Q) == 410 and len(QM) == 90, (len(Q), len(QM))


def tif(path, base, valores, codigos=None):
    """Values on the domain pixels of `base`; with `codigos` ({code: (colour, label)}) a uint8 class raster with its palette."""
    img = base.raster(valores)
    perfil = dict(driver='GTiff', width=base.w, height=base.h, count=1, crs='EPSG:4326', transform=base.aff,
                  compress='deflate', tiled=True)
    if codigos:
        perfil.update(dtype='uint8', nodata=255); img = np.where(np.isnan(img), 255, img).astype(np.uint8)
    else:
        perfil.update(dtype='float32', nodata=-9999.0); img = np.where(np.isnan(img), -9999.0, img).astype(np.float32)
    with rasterio.open(path, 'w', **perfil) as ds:
        ds.write(img, 1)
        if codigos:
            ds.write_colormap(1, {k: tuple(int(255 * x) for x in to_rgb(c)) + (255,) for k, (c, _) in codigos.items()})
            ds.update_tags(1, **{f'clase_{k}': l for k, (_, l) in codigos.items()})


def shp(path, geoms, **campos):
    datos = [np.asarray(v) for v in campos.values()]
    pyogrio.raw.write(path, np.array(shapely.to_wkb(geoms), dtype=object),
                      field_data=[v.astype(object) if v.dtype.kind in 'US' else v for v in datos], fields=list(campos), crs='EPSG:4326', geometry_type='MultiPolygon', promote_to_multi=True,
                      encoding='UTF-8')


def poligonos(g):
    """Polygon parts only (range unions made valid can carry lines or points)."""
    return shapely.multipolygons([p for p in shapely.get_parts(shapely.get_parts(g)) if p.geom_type == 'Polygon'])


def png():
    DST.mkdir(parents=True, exist_ok=True)
    for f in DST.glob('FigS*.png'):
        f.unlink()
    E.SUPL = DST
    print({g: len(v) for g, v in E.fig_atlas({g: Q[Q.grupo == g].to_dict('records') for g in GRUPOS}).items()},
          'S3.6', len(E.fig_atlas_mar(QM_PECES.to_dict('records'), archivo=MAR_FIG)),
          'S3.7', len(E.fig_atlas_aves_mar(QM_AVES.to_dict('records'), archivo=AVES_FIG)), flush=True)


def sig():
    """Class raster of every mapped species (same codes and colours as the maps) and its range polygons."""
    t, s = DST / 'sig' / 'tifs', DST / 'sig' / 'vectores'
    shutil.rmtree(t, ignore_errors=True); t.mkdir(parents=True); s.mkdir(parents=True, exist_ok=True)
    for f in s.glob('FigS*_rangos_*'):
        f.unlink()
    codigos = dict(enumerate(zip(['#ffffff', E.FUERA[0]] + [c for c, _ in E.DENTRO],
                                 ['Outside the range, with a protected analogue or protected', E.FUERA[1]] +
                                 [l.replace(E.DMIN, 'd_min') for _, l in E.DENTRO])))
    sup = E.MAR / 'superficie'
    dom_t = E.dominio_tierra()
    dom_m = E.cu.Dominio(sup, {c: E.MAR / c / 'mapas_distancia/cond100/sistema_distancia_minima.tif' for c in E.CAPAS})
    lotes = [(f"Fig{E.ATLAS_FIG[g]}_{E.G_EN[g].lower()}", Q[Q.grupo == g], E.lienzo(E.TIERRA), dom_t, ['t'],
              'BirdLife' if g == 'BIRDS' else 'UICN') for g in GRUPOS] + \
            [(MAR_FIG, QM_PECES, E.lienzo(sup), dom_m, E.CAPAS, 'UICN'),
             (AVES_FIG, QM_AVES, E.lienzo(sup), dom_m, ['superficie'], 'BirdLife')]
    for fig, esp, base, dom, capas, fuente in lotes:
        (t / fig).mkdir()
        rangos = []
        for e in esp.itertuples():
            r = E.rango(e.grupo, e.id_no); rangos.append(poligonos(r))
            for c in capas:
                img, _ = E.clases_mapa(dom, r, dom.dmin[c])
                tif(t / fig / f"{re.sub(r'[^A-Za-z0-9.-]+', '_', e.sci_name)}{'' if c == 't' else '_' + c}.tif", base, img, codigos)
        shp(s / f'{fig}_rangos_{fuente}.shp', rangos, id_no=esp.id_no.astype(int).to_numpy(),
               sci_name=esp.sci_name.to_numpy(), grupo=esp.grupo.to_numpy(), categoria=esp.categoria.to_numpy())
        print('sig', fig, len(esp), flush=True)


def docx():
    """File S3 in Word, laid out as File S2 (atlas.py): introduction, then one page per sheet of maps."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, m, Cm(2.0))
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)

    def parrafo(texto, size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6):
        """'{dmin}' becomes d (italic) with min as subscript."""
        p = doc.add_paragraph(); p.alignment = align; p.paragraph_format.space_after = Pt(after)
        for k, trozo in enumerate(texto.split('{dmin}')):
            if k:
                r = p.add_run('d'); r.italic = True; r.font.size = Pt(size)
                r = p.add_run('min'); r.font.subscript = True; r.font.size = Pt(size)
            r = p.add_run(trozo); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        return p

    n_g = Q.grupo.value_counts()
    n_m = QM_PECES.grupo.value_counts()
    t_amen, m_amen = int((Q.pct_global_en_peru <= 90).sum()), int((QM.pct_global_en_mar <= 90).sum())
    figuras = [(E.ATLAS_FIG[g], f'{E.G_EN[g]} in the continental domain ({n_g[g]} species).') for g in GRUPOS] + [
        (E.ATLAS_FIG['MARINEFISH'], f"Marine fishes ({n_m['MARINEFISH']} species), mammals ({n_m['MAMMALS']}) and reptiles "
                                     f"({n_m['REPTILES']}) in the maritime domain, in this order; sea surface left, seafloor right."),
        (E.ATLAS_FIG['SEABIRDS'], f'Seabirds at the sea surface, the only layer used for seabirds ({len(QM_AVES)} species).')]

    parrafo('File S3. Range Maps of Species in Climatic Gaps', 15, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    parrafo('Supplementary material for: Protected-area networks can represent climate worse than chance, and threatened-species '
            'plans toward 30×30 need climatic targets to close the gaps: evidence from Peru. '
            'Irwing S. Saldaña, Armando Ugaz-Cherre and Cristian Steven Sevillano-Ríos.', 10, italic=True,
            align=WD_ALIGN_PARAGRAPH.LEFT, after=14)
    parrafo(f'This file maps the {len(Q)} continental and {len(QM)} maritime species that are threatened (CR, EN or VU) or have '
            'more than 90% of their global range in Peru, or within the 200 nautical miles at sea, and that have more than half '
            'of their range in the domain in climatic gaps, in at least one layer at sea (Supplementary material S6). Of them, '
            f'{t_amen} continental and {m_amen} maritime species are threatened species with less than 90% of their global range '
            'in Peru or its sea. Figs. S3.1 to S3.5 show amphibians, mammals, reptiles, plants and birds in the continental '
            'domain, Fig. S3.6 marine fishes, mammals and reptiles at the sea surface and on the seafloor, and Fig. S3.7 '
            'seabirds at the sea surface, the only layer used for them. Within each figure, species are in alphabetical order.')
    parrafo('Within the range, dark blue marks the surface inside protected areas, light blue the unprotected surface with a '
            'protected analogue ({dmin} ≤ 0.2) and crimson the climatic gaps ({dmin} > 0.2); outside the range, pink marks the '
            'climatic gaps and white the rest of the domain. Ranges of fewer than 60 pixels are circled so that they can be '
            'seen, and maps whose window holds less than half of the domain carry an inset that locates them. The subtitle of '
            'each map gives the Red List category, the range area in the domain and its share of the global range, and the '
            'share of the range in climatic gaps.')
    parrafo('{dmin} is the Mahalanobis distance from each pixel to the most similar climate inside any protected area, computed '
            'with 11 CHELSA V2.1 predictors on land and, separately for the sea surface and the seafloor, with Bio-ORACLE v3.0 '
            'predictors at sea. Ranges come from the IUCN Red List (IUCN, 2026) and, for birds, from BirdLife International '
            'and Handbook of the Birds of the World (2025). Range maps describe the extent of occurrence and overestimate '
            'occupancy at fine resolution, so the maps show where ranges lie, not which places species occupy. The maps are '
            'derived from these ranges, which cannot be redistributed, so they are not part of the code repository.')
    total = 0
    for codigo, pie in figuras:
        paginas = sorted(DST.glob(f'Fig{codigo}_*.png'))
        assert paginas, codigo
        etiqueta = codigo.replace('-', '.')
        for k, png_ in enumerate(paginas, 1):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_picture(str(png_), width=Cm(16.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = parrafo('', 10, align=WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(f'Figure {etiqueta}. ' if k == 1 else f'Figure {etiqueta} (continued). '); r.bold = True
            r.font.size = Pt(10)
            r = p.add_run(' '.join(([pie] if k == 1 else []) + ([f'Page {k} of {len(paginas)}.'] if len(paginas) > 1 else [])))
            r.font.size = Pt(10)
        total += len(paginas)
    doc.save(DOCX)
    print('File S3', total, 'pages of maps', DOCX.stat().st_size // 2**20, 'MiB')


if __name__ == '__main__':
    if 'png' in pasos:
        png()
    if 'sig' in pasos:
        sig()
    if 'docx' in pasos:
        docx()
